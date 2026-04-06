import logging

import yaml
from openpyxl.chart import LineChart, Reference

logger = logging.getLogger(__name__)


def parse_reps(reps_value):
    """
    Parse the reps value.
    If it's a string in the form "NxM" (e.g., "5x8"), return (N, M).
    Otherwise, if it's numeric, assume (1, reps_value).
    """
    if isinstance(reps_value, str) and "x" in reps_value.lower():
        try:
            parts = reps_value.lower().split("x")
            count = int(parts[0].strip())
            reps_per_set = int(parts[1].strip())
            return count, reps_per_set
        except Exception as e:
            logger.error("Error parsing reps from '%s': %s", reps_value, e)
            return (1, 0)
    try:
        return (1, int(reps_value))
    except Exception:
        return (1, 0)


def parse_weight(weight_value):
    """
    Parse a weight like '82.5kg' or 82.5 into a float.
    """
    if weight_value is None:
        return 0.0

    try:
        return float("".join(c for c in str(weight_value) if c.isdigit() or c == "."))
    except Exception:
        logger.warning("Could not parse weight from %r", weight_value)
        return 0.0


def get_chart_data_from_yaml(yaml_path):
    """
    Loads workout data from the YAML file and computes three metrics per week:
      - Stress Index: sum of (reps × weight × %1RM)
      - Volume: sum of (reps × weight)
      - Intensity: average %1RM across all sets

    Returns:
      (weeks, stress_indices, volumes, intensities)
    """
    with open(yaml_path, "r", encoding="utf-8") as file:
        data = yaml.safe_load(file)

    if isinstance(data, dict):
        weeks_list = data.get("weeks", [])
    elif isinstance(data, list):
        weeks_list = data
    else:
        weeks_list = []

    week_numbers = []
    stress_indices = []
    volumes = []
    intensities = []
    week_index = 1

    for week_dict in weeks_list:
        for _, days in week_dict.items():
            week_numbers.append(week_index)
            week_index += 1

            set_percentages = []
            total_volume = 0.0
            total_stress = 0.0

            for day in days:
                exercises = day.get("exercises", [])
                for exercise in exercises:
                    for set_entry in exercise.get("sets", []):
                        reps_val = set_entry.get("reps")
                        weight_val = set_entry.get("weight")
                        perc_val = set_entry.get("percentage_1rm")

                        perc_float = None
                        if perc_val is not None:
                            try:
                                perc_float = float(perc_val)
                                set_percentages.append(perc_float)
                            except Exception:
                                logger.warning(
                                    "Could not parse percentage_1rm from %r", perc_val
                                )

                        if reps_val is None or weight_val is None:
                            continue

                        count, reps_per_set = parse_reps(reps_val)
                        weight_num = parse_weight(weight_val)

                        set_volume = count * reps_per_set * weight_num
                        total_volume += set_volume

                        if perc_float is not None:
                            total_stress += set_volume * (perc_float / 100.0)

            avg_intensity = (
                sum(set_percentages) / len(set_percentages) if set_percentages else 0.0
            )

            stress_indices.append(total_stress)
            volumes.append(total_volume)
            intensities.append(avg_intensity)

    return week_numbers, stress_indices, volumes, intensities


def add_charts_to_workbook(wb, yaml_path="data/program.yaml"):
    """
    Reads actual workout data from the YAML file, computes weekly metrics,
    and adds three charts to the workbook.
    """
    weeks, stress_indices, volumes, intensities = get_chart_data_from_yaml(yaml_path)

    if not weeks:
        logger.error("No week data found. Charts cannot be generated.")
        return

    if "Charts" in wb.sheetnames:
        del wb["Charts"]

    chart_sheet = wb.create_sheet("Charts")

    chart_sheet.append(["Week", "Stress Index", "Volume", "Intensity (% of 1RM)"])
    for w, stress, vol, inten in zip(weeks, stress_indices, volumes, intensities):
        chart_sheet.append([w, stress, vol, inten])

    cats = Reference(chart_sheet, min_col=1, min_row=2, max_row=len(weeks) + 1)

    stress_chart = LineChart()
    stress_chart.title = "Training Stress"
    stress_chart.y_axis.title = "Stress Index"
    stress_chart.x_axis.title = "Week"
    data_stress = Reference(chart_sheet, min_col=2, min_row=1, max_row=len(weeks) + 1)
    stress_chart.add_data(data_stress, titles_from_data=True)
    stress_chart.set_categories(cats)
    chart_sheet.add_chart(stress_chart, "F2")

    volume_chart = LineChart()
    volume_chart.title = "Volume Over Time"
    volume_chart.y_axis.title = "Volume (Reps × Weight)"
    volume_chart.x_axis.title = "Week"
    data_volume = Reference(chart_sheet, min_col=3, min_row=1, max_row=len(weeks) + 1)
    volume_chart.add_data(data_volume, titles_from_data=True)
    volume_chart.set_categories(cats)
    chart_sheet.add_chart(volume_chart, "F20")

    intensity_chart = LineChart()
    intensity_chart.title = "Intensity Progression (% of 1RM Over Time)"
    intensity_chart.y_axis.title = "% of 1RM"
    intensity_chart.x_axis.title = "Week"
    data_intensity = Reference(
        chart_sheet, min_col=4, min_row=1, max_row=len(weeks) + 1
    )
    intensity_chart.add_data(data_intensity, titles_from_data=True)
    intensity_chart.set_categories(cats)
    chart_sheet.add_chart(intensity_chart, "F38")
